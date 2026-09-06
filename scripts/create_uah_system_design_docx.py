#!/usr/bin/env python3
"""Build the UAH system design from the retained system-design template."""

from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]


def _font(size: int, *, bold: bool = False) -> ImageFont.FreeTypeFont:
    name = "arialbd.ttf" if bold else "arial.ttf"
    try:
        return ImageFont.truetype(name, size)
    except OSError:
        return ImageFont.load_default()


def _centered_text(
    draw: ImageDraw.ImageDraw,
    box: tuple[int, int, int, int],
    lines: tuple[str, ...],
    *,
    color: str,
) -> None:
    font = _font(28, bold=True)
    small = _font(20)
    heights = []
    for index, line in enumerate(lines):
        active = font if index == 0 else small
        bounds = draw.textbbox((0, 0), line, font=active)
        heights.append(bounds[3] - bounds[1])
    total = sum(heights) + 10 * (len(lines) - 1)
    y = box[1] + (box[3] - box[1] - total) / 2
    for index, line in enumerate(lines):
        active = font if index == 0 else small
        bounds = draw.textbbox((0, 0), line, font=active)
        width = bounds[2] - bounds[0]
        x = box[0] + (box[2] - box[0] - width) / 2
        draw.text((x, y), line, font=active, fill=color)
        y += heights[index] + 10


def _arrow(
    draw: ImageDraw.ImageDraw,
    start: tuple[int, int],
    end: tuple[int, int],
    *,
    color: str,
) -> None:
    draw.line((start, end), fill=color, width=6)
    x, y = end
    draw.polygon(((x, y), (x - 20, y - 12), (x - 20, y + 12)), fill=color)


def create_architecture_figure(path: Path) -> None:
    image = Image.new("RGB", (1800, 930), "#F6F8FB")
    draw = ImageDraw.Draw(image)
    navy = "#17365D"
    blue = "#2E75B6"
    pale = "#EAF2F8"
    ink = "#1F2937"
    muted = "#60758C"

    boxes = {
        "inputs": (70, 100, 380, 300),
        "compiler": (500, 100, 820, 300),
        "proposal": (940, 100, 1260, 300),
        "gate": (1380, 100, 1700, 300),
        "catalog": (1380, 530, 1700, 730),
        "owner": (940, 530, 1260, 730),
        "evidence": (500, 530, 820, 730),
        "trace": (70, 530, 380, 730),
    }
    content = {
        "inputs": ("Task + AB registry", "frame, role, required objects"),
        "compiler": ("Interaction compiler", "minimal task-scoped closure"),
        "proposal": ("Model proposal", "typed operation; no effect truth"),
        "gate": ("Deterministic gate", "role, reach, policy, minefields"),
        "catalog": ("Binding catalog", "candidate quarantine + mode"),
        "owner": ("Environment owner", "approved AB1 implementation"),
        "evidence": ("Effect evidence", "owner-issued observables"),
        "trace": ("Trace + evaluation", "replay, promotion, rollback"),
    }
    for key, box in boxes.items():
        fill = navy if key in {"gate", "owner"} else pale
        outline = navy if key in {"gate", "owner"} else blue
        text_color = "white" if key in {"gate", "owner"} else ink
        draw.rounded_rectangle(box, radius=22, fill=fill, outline=outline, width=5)
        _centered_text(draw, box, content[key], color=text_color)

    for left, right in (("inputs", "compiler"), ("compiler", "proposal"), ("proposal", "gate")):
        a, b = boxes[left], boxes[right]
        _arrow(draw, (a[2], (a[1] + a[3]) // 2), (b[0], (b[1] + b[3]) // 2), color=blue)
    _arrow(draw, (1540, 300), (1540, 530), color=blue)
    for left, right in (("catalog", "owner"), ("owner", "evidence"), ("evidence", "trace")):
        a, b = boxes[left], boxes[right]
        _arrow(draw, (a[0], (a[1] + a[3]) // 2), (b[2], (b[1] + b[3]) // 2), color=blue)
    draw.text(
        (70, 825),
        "Invariant: semantic identity is stable; bindings are replaceable; only owners close effects.",
        font=_font(25, bold=True),
        fill=muted,
    )
    path.parent.mkdir(parents=True, exist_ok=True)
    image.save(path)


def _paragraph(document: Document, index: int, text: str) -> None:
    paragraph = document.paragraphs[index]
    paragraph.clear()
    paragraph.add_run(text)


def _cell(table, row: int, column: int, text: str) -> None:
    table.cell(row, column).text = text


def populate(reference: Path, output: Path, figure: Path) -> None:
    document = Document(reference)
    create_architecture_figure(figure)

    _paragraph(document, 8, "Universal Agentic Harness")
    _paragraph(document, 9, "Shadow-First Semantic Binding Architecture")

    metadata = document.tables[0]
    _cell(metadata, 0, 0, "STATUS\nProposed · H1 slice active")
    _cell(metadata, 0, 2, "OWNER\nNeural Workbench / UAH")
    _cell(metadata, 0, 4, "LAST UPDATED\nAugust 4, 2026")

    identity = document.tables[1]
    values = (
        ("Authors", "UAH research and implementation track"),
        ("Reviewers", "NAO owners; Neural Workbench owner; runtime/eval owner"),
        (
            "Related docs",
            "universal_agentic_harness_masterplan.md; "
            "universal_agentic_harness_development_log.md; ADR 0001",
        ),
        ("Decision status", "Semantic binding boundary accepted; live coupling pending"),
    )
    for row, (label, value) in enumerate(values):
        _cell(identity, row, 0, label)
        _cell(identity, row, 1, value)

    _paragraph(
        document,
        22,
        "UAH is a portable control kernel that scopes a model to a frame-relative "
        "AB projection, accepts only typed proposals, and delegates execution and "
        "effect proof to environment owners. The first implementation mounts "
        "recorded NAO-shaped chatbot and planner outputs over an in-process fake "
        "owner. It proves the semantic seam without importing ROS, NAO packages, "
        "provider SDKs, or live runtime products.",
    )
    _paragraph(
        document,
        23,
        "The design supports local and frontier models through later adapters, "
        "but this proposal does not claim live-node parity, Bonsai/Watson parity, "
        "safe online learning, H1 completion, or global meaning for an AB level.",
    )

    goals = document.tables[2]
    goal_rows = (
        ("Goals", "Non-goals"),
        ("Stable semantic objects across environment changes", "Mirror every private method in the AB registry"),
        ("Fail-closed proposal, binding, and effect gates", "Let discovery grant authority or prove effects"),
        ("Replayable model-harness-environment evaluation", "Replace NAO package ownership"),
        ("Quarantined, reversible Workbench candidates", "Online publication of learned code or AB objects"),
    )
    for row, values in enumerate(goal_rows):
        for column, value in enumerate(values):
            _cell(goals, row, column, value)

    _paragraph(
        document,
        28,
        "NAO already contains mature dialogue, planner, orchestrator, perception, "
        "and fake-skill seams. ZeroTier experiments show that Bonsai is materially "
        "more memory-efficient than the Watson control, but its strict routed "
        "workflow is not yet equivalent. Selecting a model by memory or throughput "
        "alone would confound model, harness, runtime, adapter, and evaluator.",
    )
    _paragraph(
        document,
        29,
        "The system therefore governs one immutable configuration at a time. "
        "The UAH owns projection, typed gates, trace grammar, evaluation, and "
        "promotion governance. Existing environment components retain policy, "
        "execution, observations, and effect ownership.",
    )

    picture = document.paragraphs[31]
    picture.clear()
    picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture.add_run().add_picture(str(figure), width=Inches(6.75))
    _paragraph(document, 33, "Figure 1. UAH semantic binding and evidence lifecycle.")

    components = document.tables[3]
    component_rows = (
        ("Component", "Responsibility", "Primary storage", "Failure behavior"),
        ("AB registry snapshot", "Semantic objects, levels, owners, decomposition", "Versioned JSON + SHA-256", "Unknown objects fail closed"),
        ("Interaction compiler", "Build minimal role/task closure", "Immutable module spec", "Out-of-band object rejected"),
        ("Output gate", "Validate type, reach, level, claims", "Decision in trace", "No dispatch on rejection"),
        ("Binding catalog", "Map semantics to reviewed environment APIs", "Revisioned bindings", "Candidate, missing, or ambiguous binding stops"),
        ("Environment owner", "Execute approved AB1 and issue observations", "Owner runtime", "Failure remains owner-attributed"),
    )
    for row, values in enumerate(component_rows):
        for column, value in enumerate(values):
            _cell(components, row, column, value)

    lifecycle = (
        "1. Compile — receive the task, role, frame, registry identity, and required objects.",
        "2. Project — close the minimal inspectable decomposition and direct-control band.",
        "3. Propose — a model or recorded replay emits a typed operation without effect truth.",
        "4. Gate — validate output ownership, projection reach, permissions, and minefields.",
        "5. Resolve — select one approved binding for the environment and runtime mode.",
        "6. Dispatch — the environment owner executes under budgets, cancellation, and retries.",
        "7. Verify — close terminal observables from owner evidence and append lifecycle events.",
    )
    for index, text in enumerate(lifecycle, start=40):
        _paragraph(document, index, text)

    contract = document.tables[4]
    contract_rows = (
        ("Field", "Type", "Required", "Description"),
        ("object_id", "string", "Yes", "Stable frame-relative AB semantic identity"),
        ("binding_id", "string", "Yes", "Versioned implementation pointer identity"),
        ("implementation_owner", "string", "Yes", "Package/component that implements this pointer"),
        ("source_revision", "string", "Yes", "Exact source revision used for validation"),
        ("runtime_modes", "string[]", "Yes", "Modes in which the binding may resolve"),
        ("status", "enum", "Yes", "candidate, approved, or disabled"),
        ("evidence_ref", "string", "For execution", "Durable owner-issued evidence reference"),
    )
    for row, values in enumerate(contract_rows):
        for column, value in enumerate(values):
            _cell(contract, row, column, value)

    guarantees = (
        "Only approved bindings resolve; discovery and declaration produce candidates.",
        "AB0 interfaces may have producer, contract, and consumer pointers but cannot dispatch.",
        "An executable AB1 binding must be implemented by the registry-declared effect owner.",
        "Every result captures registry, binding, environment, source, and configuration identity.",
        "Model text, runtime discovery, and successful transport are never effect truth.",
        "Canonical schemas and implementation details are versioned independently.",
    )
    for index, text in enumerate(guarantees, start=54):
        _paragraph(document, index, text)

    _paragraph(
        document,
        63,
        "Replay replaces the model with recorded proposals while preserving the "
        "registry, task, bindings, environment, and evaluator. Binding identity "
        "and evidence references make duplicate attempts visible. Runtime owners "
        "remain responsible for idempotency; the UAH records attempt lineage and "
        "never treats a repeated response as a new effect without fresh evidence.",
    )
    consistency = document.tables[5]
    consistency_rows = (
        ("Scenario", "Expected behavior", "Reasoning"),
        ("Duplicate proposal", "Record a new attempt; owner idempotency applies", "Proposal identity is not effect identity"),
        ("Candidate or ambiguous binding", "Fail before dispatch", "Discovery cannot become authority"),
        ("Owner execution failure", "Retain failure evidence; terminal remains open", "Transport success is not task success"),
        ("Stale or missing evidence", "Reject completion and permit bounded recovery", "Fresh owner proof closes effects"),
    )
    for row, values in enumerate(consistency_rows):
        for column, value in enumerate(values):
            _cell(consistency, row, column, value)

    security = (
        "Authorization is compiled from role, frame, task, registry owner, and approved binding state.",
        "Trace payloads minimize sensitive model, tool, and environment data; artifacts use references.",
        "Provider credentials and runtime secrets remain outside the semantic registry and prompt.",
        "Replay, promotion, and debugging tools default to read-only and shadow execution.",
        "Adaptive structures remain quarantined until counterexamples, holdout, owner review, provenance, and rollback pass.",
    )
    for index, text in enumerate(security, start=67):
        _paragraph(document, index, text)

    operations = document.tables[6]
    operation_rows = (
        ("Signal", "SLO or alert", "Owner", "Launch gate"),
        ("Terminal effect closure", "100% for deterministic canaries", "Environment + evaluator", "Required"),
        ("Scope violations", "0 admitted out-of-projection actions", "UAH gate", "Required"),
        ("Trace reconstruction", "100% deterministic replay agreement", "UAH trace", "Required"),
        ("Resource headroom", "Configured minimum RAM/VRAM remains free", "Runtime owner", "Required"),
        ("Latency", "Within role-specific TTFT and terminal budget", "Provider/runtime", "Required"),
        ("Rollback", "Promoted configuration can be reverted and replayed", "Promotion owner", "Required"),
    )
    for row, values in enumerate(operation_rows):
        for column, value in enumerate(values):
            _cell(operations, row, column, value)

    alternatives = document.tables[7]
    alternative_rows = (
        ("Alternative", "Why it was considered", "Why it was not selected"),
        ("Promote every method into AB", "Automatic environment ingestion", "Private refactors churn semantics and discovery becomes authority"),
        ("Import NAO/ROS into core", "Fast reuse of mature nodes", "Breaks portability and ownership boundaries"),
        ("Begin with live autonomous coupling", "Earlier end-to-end behavior", "No parity, attribution, or rollback evidence"),
        ("Fork a broad agent product", "Ready providers and sessions", "Runtime policy would dominate the AB semantic kernel"),
    )
    for row, values in enumerate(alternative_rows):
        for column, value in enumerate(values):
            _cell(alternatives, row, column, value)

    questions = (
        "1. What freshness and clock contract must every effect-evidence type carry?",
        "2. Which frozen NAO fixtures form the first held-out shadow-parity suite?",
        "3. What provider-neutral request/result protocol is sufficient for Watson and Bonsai?",
        "4. Which owner signs a candidate binding and which independent evaluator can reject it?",
    )
    for index, text in enumerate(questions, start=81):
        _paragraph(document, index, text)

    _paragraph(
        document,
        87,
        "Approve the semantic-object/implementation-binding boundary and the "
        "H2 v0 qualification contract. Continue shadow-first. Next, serialize "
        "the complete configuration and TaskSpec, "
        "append lifecycle events for the recorded success and a stale-evidence "
        "counterexample, and require replay-identical terminal decisions. Only "
        "then add the Watson/Bonsai provider adapter and NAO shadow parity.",
    )
    milestones = document.tables[8]
    milestone_rows = (
        ("Milestone", "Deliverable", "Exit criteria"),
        ("M1", "Semantic bindings and fake-owner qualification", "Implemented; portable tests green"),
        ("M2", "Configuration/TaskSpec lifecycle replay", "Success and stale failure replay identically"),
        ("M3", "Watson/Bonsai ROS-free eval runner", "Frozen repeated matrix with effect graders"),
        ("M4", "NAO read-only shadow adapter", "Package-owned parity and disagreement ledger"),
    )
    for row, values in enumerate(milestone_rows):
        for column, value in enumerate(values):
            _cell(milestones, row, column, value)

    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--reference",
        type=Path,
        default=ROOT / "docs/artifacts/system-design/system-design-reference.docx",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=ROOT / "docs/artifacts/system-design/universal_agentic_harness_system_design.docx",
    )
    parser.add_argument(
        "--figure",
        type=Path,
        default=ROOT / "docs/artifacts/system-design/uah-system-architecture.png",
    )
    args = parser.parse_args()
    populate(args.reference, args.output, args.figure)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
