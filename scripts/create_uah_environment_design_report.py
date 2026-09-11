#!/usr/bin/env python3
"""Build the UAH environment design report from the retained design template."""

from __future__ import annotations

import argparse
from pathlib import Path
import shutil
import textwrap
import zipfile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_REFERENCE = Path(
    "/home/juanbeck/.codex/plugins/cache/openai-curated-remote/"
    "openai-templates/0.1.1/skills/artifact-template-design-report/"
    "assets/reference.docx"
)
DEFAULT_OUTPUT_DIR = ROOT / "docs" / "artifacts" / "design-reports"


def _font(size: int, *, bold: bool = False) -> ImageFont.FreeTypeFont:
    filename = "DejaVuSans-Bold.ttf" if bold else "DejaVuSans.ttf"
    try:
        return ImageFont.truetype(filename, size)
    except OSError:
        return ImageFont.load_default()


def _centered_lines(
    draw: ImageDraw.ImageDraw,
    box: tuple[int, int, int, int],
    title: str,
    subtitle: str,
    *,
    title_color: str = "#17233D",
    subtitle_color: str = "#44546A",
) -> None:
    title_font = _font(30, bold=True)
    subtitle_font = _font(21)
    lines: list[tuple[str, ImageFont.FreeTypeFont, str]] = []
    lines.extend((line, title_font, title_color) for line in textwrap.wrap(title, 24))
    lines.extend(
        (line, subtitle_font, subtitle_color) for line in textwrap.wrap(subtitle, 34)
    )
    heights = [draw.textbbox((0, 0), line, font=font)[3] for line, font, _ in lines]
    total = sum(heights) + 9 * max(0, len(lines) - 1)
    y = box[1] + (box[3] - box[1] - total) / 2
    for (line, font, color), height in zip(lines, heights, strict=True):
        bounds = draw.textbbox((0, 0), line, font=font)
        x = box[0] + (box[2] - box[0] - (bounds[2] - bounds[0])) / 2
        draw.text((x, y), line, font=font, fill=color)
        y += height + 9


def _box(
    draw: ImageDraw.ImageDraw,
    bounds: tuple[int, int, int, int],
    title: str,
    subtitle: str,
    *,
    fill: str,
    outline: str,
) -> None:
    draw.rounded_rectangle(bounds, radius=25, fill=fill, outline=outline, width=5)
    _centered_lines(draw, bounds, title, subtitle)


def _arrow(
    draw: ImageDraw.ImageDraw,
    start: tuple[int, int],
    end: tuple[int, int],
    *,
    label: str | None = None,
    color: str = "#8BA1C0",
) -> None:
    draw.line((start, end), fill=color, width=7)
    x, y = end
    draw.polygon(((x, y), (x - 17, y - 12), (x + 17, y - 12)), fill=color)
    if label:
        font = _font(18, bold=True)
        bounds = draw.textbbox((0, 0), label, font=font)
        tx = (start[0] + end[0] - (bounds[2] - bounds[0])) / 2
        ty = (start[1] + end[1]) / 2 - 28
        draw.rounded_rectangle(
            (tx - 8, ty - 4, tx + bounds[2] - bounds[0] + 8, ty + bounds[3] + 4),
            radius=8,
            fill="#17233D",
        )
        draw.text((tx, ty), label, font=font, fill="#FFFFFF")


def create_cover_figure(path: Path) -> None:
    image = Image.new("RGB", (1800, 1800), "#101A2F")
    draw = ImageDraw.Draw(image)
    draw.text((95, 65), "UAH RUNTIME SPINE", font=_font(50, bold=True), fill="#FFFFFF")
    draw.text(
        (95, 135),
        "Identity, environment, causal trace, and evidence closure",
        font=_font(26),
        fill="#AFC7E8",
    )

    boxes = {
        "role": (95, 250, 750, 445),
        "profile": (1050, 250, 1705, 445),
        "agent": (95, 565, 750, 760),
        "environment": (1050, 565, 1705, 760),
        "runs": (300, 880, 1500, 1075),
        "trace": (300, 1195, 1500, 1390),
        "acceptance": (300, 1510, 1500, 1705),
    }
    content = {
        "role": (
            "role_configuration_id",
            "Frozen frame, capabilities, authority, and admission profile",
        ),
        "profile": (
            "environment_profile_id",
            "Domain contract, interfaces, readiness, and stable agent roster",
        ),
        "agent": (
            "agent_id + stable handle",
            "Role, model, prompt, harness, adapters, and fidelity revision",
        ),
        "environment": (
            "environment_run_id",
            "One owner-attested native activation and ingress envelope",
        ),
        "runs": (
            "attached agent_run_id actors",
            "Persistent across tasks and model leases; standby releases hardware",
        ),
        "trace": (
            "task_id -> trace_id -> operation graph",
            "One causal record; actor views are read-only projections",
        ),
        "acceptance": (
            "EffectObligation -> TaskAcceptance",
            "Required effects decide closure; best-effort deficits remain visible",
        ),
    }
    fills = {
        "role": "#E9EDFF",
        "profile": "#DFF5F4",
        "agent": "#E9EDFF",
        "environment": "#DFF5F4",
        "runs": "#EDE6FF",
        "trace": "#DFF5F4",
        "acceptance": "#FFF2CC",
    }
    outlines = {
        "role": "#536DFE",
        "profile": "#12939A",
        "agent": "#536DFE",
        "environment": "#12939A",
        "runs": "#8E5AD7",
        "trace": "#12939A",
        "acceptance": "#C58A00",
    }
    for key, bounds in boxes.items():
        _box(draw, bounds, *content[key], fill=fills[key], outline=outlines[key])

    _arrow(draw, (423, 445), (423, 565))
    _arrow(draw, (1378, 445), (1378, 565))
    _arrow(draw, (423, 760), (650, 880))
    _arrow(draw, (1378, 760), (1150, 880))
    _arrow(draw, (900, 1075), (900, 1195), label="actor events")
    _arrow(draw, (900, 1390), (900, 1510), label="owner evidence")
    path.parent.mkdir(parents=True, exist_ok=True)
    image.save(path)


def create_report_result_figure(path: Path) -> None:
    image = Image.new("RGB", (1800, 980), "#F5F7FB")
    draw = ImageDraw.Draw(image)
    draw.text(
        (80, 45),
        "NAO report_result: one AB1 operation, explicit delegation",
        font=_font(38, bold=True),
        fill="#17233D",
    )
    boxes = {
        "report": (70, 190, 420, 380),
        "verify": (520, 80, 880, 270),
        "compose": (520, 370, 880, 560),
        "invoke": (980, 370, 1325, 560),
        "artifact": (1425, 370, 1770, 560),
        "deliver": (980, 690, 1325, 880),
        "evidence": (1425, 690, 1770, 880),
    }
    content = {
        "report": ("report_result AB1", "Planner-visible object in nao_runtime"),
        "verify": ("verify evidence", "Same-frame deterministic owner"),
        "compose": ("compose grounded report", "Delegated to chatbot in nao_dialogue"),
        "invoke": ("model invocation", "Bounded prompt under chatbot run"),
        "artifact": ("GroundedReportArtifact", "Typed text with evidence references"),
        "deliver": ("deliver report", "Native communication owner"),
        "evidence": ("speech evidence", "Required or best-effort obligation"),
    }
    palette = {
        "report": ("#DFF5F4", "#12939A"),
        "verify": ("#EDE6FF", "#8E5AD7"),
        "compose": ("#EDE6FF", "#8E5AD7"),
        "invoke": ("#E9EDFF", "#536DFE"),
        "artifact": ("#FFF2CC", "#C58A00"),
        "deliver": ("#FDE6E1", "#D75A4A"),
        "evidence": ("#FFF2CC", "#C58A00"),
    }
    for key, bounds in boxes.items():
        _box(draw, bounds, *content[key], fill=palette[key][0], outline=palette[key][1])

    _arrow(draw, (420, 260), (520, 180), label="decomposes_to")
    _arrow(draw, (420, 320), (520, 465), label="delegates_to")
    _arrow(draw, (880, 465), (980, 465))
    _arrow(draw, (1325, 465), (1425, 465))
    _arrow(draw, (1598, 560), (1150, 690), label="typed artifact")
    _arrow(draw, (1325, 785), (1425, 785))
    path.parent.mkdir(parents=True, exist_ok=True)
    image.save(path)


def _replace_paragraph(document: Document, index: int, text: str) -> None:
    paragraph = document.paragraphs[index]
    paragraph.clear()
    paragraph.add_run(text)


def _replace_cover_image(docx_path: Path, cover_path: Path) -> None:
    replacement = cover_path.read_bytes()
    temporary = docx_path.with_suffix(".zip")
    with (
        zipfile.ZipFile(docx_path, "r") as source,
        zipfile.ZipFile(temporary, "w", zipfile.ZIP_DEFLATED) as target,
    ):
        for item in source.infolist():
            payload = (
                replacement
                if item.filename == "word/media/image1.png"
                else source.read(item.filename)
            )
            if item.filename == "word/header1.xml":
                payload = payload.replace(
                    b"Report title", b"UAH architecture checkpoint"
                )
                payload = payload.replace(b">Date<", b">9 September 2026<")
            if item.filename == "word/document.xml":
                payload = payload.replace(
                    b">Introduction<", b">Architecture checkpoint<"
                )
                payload = payload.replace(
                    b">Notes<", b">Identity and authority map<"
                )
                payload = payload.replace(
                    b">Source placeholders<", b">Evidence sources<"
                )
            target.writestr(item, payload)
    temporary.replace(docx_path)


def populate(reference: Path, output_dir: Path) -> Path:
    output_dir.mkdir(parents=True, exist_ok=True)
    retained_reference = output_dir / "design-report-reference.docx"
    shutil.copy2(reference, retained_reference)

    cover = output_dir / "uah-runtime-spine.png"
    report_result = output_dir / "uah-nao-report-result-delegation.png"
    create_cover_figure(cover)
    create_report_result_figure(report_result)

    document = Document(reference)
    document.core_properties.title = (
        "Universal Agentic Harness Environment Design Report"
    )
    document.core_properties.subject = (
        "H0-H2 environment, ingress, identity, trace, and task-closure architecture"
    )
    document.core_properties.author = "Universal Agentic Harness project"

    replacements = {
        2: "Universal Agentic Harness",
        4: "Contents",
        5: "",
        6: "Executive summary",
        7: "The H0-H2 architecture now has an explicit runtime envelope. An owner-attested EnvironmentRun groups native ingress, persistent agent actors, domain tasks, causal traces, and evidence. Agent identity remains independent of provider placement, while each model invocation records the exact lease and loaded instance that supplied computation.",
        8: "Task closure is no longer inferred from planner status, model text, or an event label. Tasks declare required and best-effort effect obligations before execution. A deterministic evaluator compares those obligations with owner-issued evidence and produces accepted, accepted-with-deficit, suspended, or rejected outcomes.",
        9: "At a glance",
        10: "H0 freezes identity, environment, operation, acceptance, and trace contracts. It adds no scheduler.",
        11: "H1 adds lifecycle registries, ingress, standby, fixed leases, PromptCompiler, admission, replay, O1, and VerifiedTraceDigest.",
        12: "H2 is the first NAO demonstration: fixed chatbot and planner handles, native lifecycle ownership, and report_result as AB1 with explicit chatbot delegation.",
        13: "Architecture checkpoint",
        14: "The identity hierarchy distinguishes reusable role configuration, immutable agent embodiment, stable handle revision, bounded agent activation, provider capacity, and individual model invocation. Registering an agent does not load or call a model. An active agent run may release its lease and remain in standby inside one environment activation.",
        15: "EnvironmentIngress separates native stimuli from task semantics. Deterministic TaskIngressPolicy decides whether an item updates environment state, starts work, resumes work, notifies an existing task, or is rejected. A model can interpret admitted content but cannot rewrite task or trace lineage.",
        16: "Key findings",
        17: "The accepted design preserves the portable-kernel boundary. Core contracts contain no ROS, NAO, provider-SDK, or runtime-product imports. NAO compatibility and recorded qualification now reside in the separate ab_harness_nao package, with a regression test enforcing dependency direction. DomainContractPack carries frame-relative objects, role projections, bindings, environment profiles, evidence rules, prompt policy, and qualification cases. Custom adapters are limited to irreducible native normalization.",
        18: "Context and conditions",
        19: "The source-resolution audit used NAO tag v1.0.0 at ebffe93, chatbot revision a2ecca796, and intended NeuralWorkbench revision e76ba7e. Focused read-only checks passed 112 chatbot turn-engine cases and 41 planner supervisor/gate cases. These baselines constrain compatibility but do not qualify UAH H2.",
        20: "Patterns in the evidence",
        21: "NAO already contains mature routing, goal and plan lifecycle, cancellation, supersession, stale-version checks, owner dispatch, feedback, and report_result behavior. UAH should normalize and trace these seams rather than copy their domain policy. The current NeuralWorkbench report_result decomposition is stale relative to the authoritative NAO runtime and requires an owner-reviewed DomainContractPack revision.",
        23: "Key takeaway. One causal trace may include several agent runs, but every event names its actor and every operation retains one frame-relative AB coordinate. Cross-frame work is visible as delegation, never hidden as ordinary decomposition.",
        25: "Implications",
        26: "Observatory and NeuralWorkbench remain separate. Observatory renders immutable facts and projections. Workbench consumes verified trace digests to retrieve experience and compare candidates, but it cannot mutate the ledger or decide when evidence becomes memory or authority. H2 provides only the attachment seams; retrieval and model-assisted search begin in H3.",
        27: "Recommendations",
        28: "Implementation should proceed through small public contracts and preserve the current H0 compatibility surface while the new identity grammar lands.",
        29: "Retain the implemented TaskAcceptanceEvaluator.evaluate(effect_obligations, evidence_set) -> TaskAcceptance boundary and compile obligations only from frozen task and domain contracts.",
        30: "Retain profile-verified EnvironmentRun registration, then classify one immutable ingress, attach one standby-capable agent run, and connect the existing typed proposal, fake-owner, and acceptance path.",
        31: "Produce a replayable VerifiedTraceDigest, then freeze the NAO DomainContractPack and test planner parity, report_result delegation, required-effect failure, and best-effort deficit paths.",
        32: "Conclusion",
        33: "The architecture grill is closed. The public task-acceptance and profile-verified EnvironmentRun registration seams are implemented. The design supports continuous environment-bound agents without permanent model allocation, preserves domain authority, and creates a trace grammar suitable for O1/O2 debugging and later Workbench memory.",
        34: "Environment profiles and run attestations are immutable in-memory contracts without persistence or signature verification. No TaskSpec obligation compiler, ingress classifier, multi-actor ledger, report_result adapter, or dynamic allocator is implemented at this checkpoint. The current release remains an H0 contract proof with partial H1 vertical slices.",
        35: "Appendix",
        36: "Identity and authority map",
        37: "Semantic chain: DomainContractPack -> AgentRoleConfiguration -> AgentManifest -> agent_id -> AgentHandleRevision -> agent_run_id.",
        38: "Runtime chain: EnvironmentProfile -> environment_run_id -> EnvironmentIngress -> TaskIngressDecision -> task_id -> trace_id -> operation graph -> effect obligations -> TaskAcceptance.",
        39: "Evidence sources",
        40: "Universal Agentic Harness foundation, canonical Markdown, 9 September 2026.",
        41: "Universal Agentic Harness masterplan and development log, canonical Markdown, 9 September 2026.",
        42: "NAO ROS4HRI bridge, annotated tag v1.0.0, peeled commit ebffe93a74be4e013ce0f60fdfc41268dba73fc3.",
        43: "Chatbot planner-hook source, revision a2ecca796; NeuralWorkbench intended companion revision e76ba7e.",
        45: "Report status. The architecture grill is closed. Task acceptance and profile-verified environment registration are implemented. Ingress, lifecycle replay, and H2 qualification remain open.",
    }
    for index, value in replacements.items():
        _replace_paragraph(document, index, value)

    checkpoint_heading = document.paragraphs[13].paragraph_format
    checkpoint_heading.space_before = Pt(18)
    checkpoint_heading.keep_with_next = True

    picture = document.paragraphs[24]
    picture.clear()
    picture.style = document.styles["normal"]
    picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
    inline_shape = picture.add_run().add_picture(str(report_result), width=Inches(6.35))
    inline_shape._inline.docPr.set(
        "descr",
        "NAO report_result AB1 operation with same-frame evidence verification, "
        "cross-frame chatbot delegation, and native communication evidence.",
    )
    inline_shape._inline.docPr.set("title", "NAO report_result delegation")

    metadata = document.tables[0]
    metadata.cell(0, 0).text = "Environment, ingress, and task closure architecture"
    metadata.cell(0, 2).text = "Prepared by UAH project\n9 September 2026"

    findings = document.tables[1]
    rows = (
        ("Theme", "Observation", "Implication"),
        (
            "Identity",
            "Logical agents and provider resources have separate lifetimes",
            "Hardware changes cannot silently rewrite agent provenance",
        ),
        (
            "Environment",
            "One attested activation groups ingress, actors, tasks, traces, and evidence",
            "Container restarts and cross-agent workflows remain distinguishable",
        ),
        (
            "Closure",
            "Required and best-effort effects are evaluated against owner evidence",
            "Successful operations survive later reporting deficits",
        ),
    )
    for row_index, values in enumerate(rows):
        for column_index, value in enumerate(values):
            findings.cell(row_index, column_index).text = value

    output = output_dir / "uah_environment_ingress_task_closure_design_report.docx"
    document.save(output)
    _replace_cover_image(output, cover)
    return output


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--reference", type=Path, default=DEFAULT_REFERENCE)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)
    args = parser.parse_args()
    output = populate(args.reference, args.output_dir)
    print(output)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
